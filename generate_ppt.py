"""
Markdown → PowerPoint (.pptx) 変換スクリプト
使い方:
  python generate_ppt.py input.md output.pptx [logo.png]

Markdown記法:
  # タイトル                          → タイトルスライド
  サブタイトル or メタ情報テキスト     → タイトルのサブタイトルに追加
  ### セクション名                     → セクション区切りスライド（ダークネイビー全面）
  ## 1.1. 節番号 | スライドタイトル   → コンテンツスライド（節番号付き）
  ## スライドタイトル                  → コンテンツスライド（節番号なし）
  > リード文                           → タイトル直下の概要メッセージ
  - 箇条書き                           → 標準箇条書き（●）
    - サブ項目                         → 字下げ箇条書き（–）
  1. 番号付きリスト                    → 番号付きリスト
  ---                                  → コンテンツ内の視覚的区切り
  | ヘッダー1 | ヘッダー2 |           → テーブル（Markdown標準記法）
  |---|---|                            → テーブルヘッダー区切り行
  ::left [タイトル]                   → 2カラムレイアウト（左カラム開始）
  ::right [タイトル]                  → 2カラムレイアウト（右カラム開始）
  :::                                 → カラムレイアウト終了
"""

import sys
import re
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree


# ===== SALESCORE カラーパレット（salescore-pptx-format ベース）=====
COLOR_PRIMARY    = RGBColor(0x1A, 0x1A, 0x1A)   # #1a1a1a 基本テキスト・ライン
COLOR_ACCENT     = RGBColor(0x00, 0x66, 0xCD)   # #0066cd アクセント（記号・強調）
COLOR_DANGER     = RGBColor(0xE7, 0x4C, 0x3B)   # #e74c3b 重要ポイント（ごくわずか）
COLOR_MUTED      = RGBColor(0x88, 0x88, 0x88)   # #888888 フッター・ページ番号
COLOR_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)   # 背景・テーブルデータ行
COLOR_LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)   # 薄グレー背景
COLOR_MID_GRAY   = RGBColor(0xCC, 0xCC, 0xCC)   # 区切り線・縦区切り
COLOR_STRIPE     = RGBColor(0xF5, 0xF7, 0xFA)   # テーブル縞模様
# セクション・タイトルスライド用（補足レイアウト）
COLOR_DARK_NAVY  = RGBColor(0x1D, 0x34, 0x61)   # セクション全面・タイトル背景
COLOR_NAVY       = RGBColor(0x1B, 0x4F, 0xA8)   # セクションアクセント
COLOR_BLUE       = RGBColor(0x1E, 0x6F, 0xD8)   # セクションアクセントライン
# 後方互換エイリアス
COLOR_TEXT       = COLOR_PRIMARY
COLOR_BLACK      = COLOR_PRIMARY
COLOR_GRAY       = COLOR_MUTED
COLOR_GRAY_LIGHT = COLOR_MID_GRAY
COLOR_SECTION_NUM = COLOR_MUTED

# フォント（salescore-pptx-format 必須）
FONT_NAME = "Meiryo UI"

# スライドサイズ（16:9）
SLIDE_WIDTH  = Inches(13.33)
SLIDE_HEIGHT = Inches(7.5)

# ===== 共通レイアウト座標（salescore-pptx-format 実測値）=====

# ロゴ（右上）
LOGO_X = Inches(10.213)
LOGO_Y = Inches(0.091)
LOGO_W = Inches(2.205)
LOGO_H = Inches(0.512)

# コンテンツ左端・幅
CONTENT_X  = Inches(0.917)
CONTENT_RW = Inches(11.5)

# タイトルエリア
TITLE_AREA_X = Inches(0.906)
TITLE_AREA_W = Inches(9.5)
TITLE_AREA_Y = Inches(0.276)
TITLE_AREA_H = Inches(0.315)

# ヘッダーライン
HEADER_LINE_Y = Inches(0.689)
HEADER_LINE_H = Pt(1.5)

# リード文エリア
LEAD_Y = Inches(0.700)
LEAD_H = Inches(0.688)

# コンテンツエリア開始Y
CONTENT_AREA_Y = Inches(1.4)

# フッターライン
FOOTER_LINE_Y = Inches(6.951)
FOOTER_LINE_H = Pt(1.5)

# フッターテキスト・ページ番号
FOOTER_TEXT_Y = Inches(7.0)
FOOTER_TEXT_H = Inches(0.25)

FOOTER_TEXT = "Confidential All Rights Reserved SALESCORE Co., Ltd."

# ===== フォントサイズ階層（salescore-pptx-format 準拠）=====
FONT_SLIDE_TITLE = 21   # スライドタイトル（bold）
FONT_LEAD        = 18   # リードテキスト
FONT_BODY        = 16   # ボディ標準（箇条書き）
FONT_BODY_SUB    = 14   # ボディ補足（サブ箇条書き・テーブル）
FONT_FOOTER      = 9    # フッターテキスト
FONT_PAGE_NUM    = 12   # ページ番号

# テーブル行高さ
TABLE_HEADER_ROW_H = Inches(0.40)
TABLE_DATA_ROW_H   = Inches(0.35)


# ─────────────────────────────────────────────
#  ユーティリティ
# ─────────────────────────────────────────────

def set_font(run, size, bold=False, color=None, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def add_inline_text(para, text, size, default_bold=False, default_color=None):
    """**bold**マークダウンを解析して書式付きRunに変換する"""
    if not text:
        return
    parts = re.split(r'\*\*(.+?)\*\*', text)
    for i, part in enumerate(parts):
        if not part:
            continue
        r = para.add_run()
        r.text = part
        is_bold = default_bold or (i % 2 == 1)  # 奇数インデックス = **...**の中身
        set_font(r, size, bold=is_bold, color=default_color)


def remove_shadow(shape):
    """シェイプ影エフェクトを除去（テーマ継承を上書き）"""
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    etree.SubElement(spPr, qn('a:effectLst'))


def add_rect(slide, x, y, w, h, color, line_color=None, line_w=Pt(0)):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_w if line_w else Pt(0.75)
    else:
        shape.line.fill.background()
    remove_shadow(shape)
    return shape


def set_cell_bg(cell, color):
    """テーブルセルの背景色をXMLで直接設定"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in (qn('a:solidFill'), qn('a:noFill'), qn('a:gradFill'), qn('a:pattFill')):
        for el in tcPr.findall(tag):
            tcPr.remove(el)
    sf = etree.SubElement(tcPr, qn('a:solidFill'))
    clr = etree.SubElement(sf, qn('a:srgbClr'))
    clr.set('val', str(color))


def add_logo(slide, logo_path):
    if logo_path and os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, LOGO_X, LOGO_Y, width=LOGO_W)
    else:
        # フォールバック: テキスト代替
        tb = slide.shapes.add_textbox(LOGO_X, LOGO_Y, LOGO_W, LOGO_H)
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "SALESCORE"
        set_font(r, 16, bold=True, color=COLOR_PRIMARY)


def add_header_line(slide):
    """ヘッダーライン（salescore-pptx-format 実測値）"""
    add_rect(slide, CONTENT_X, HEADER_LINE_Y, CONTENT_RW, HEADER_LINE_H, COLOR_PRIMARY)


def add_footer(slide, page_num=None):
    """フッターライン + テキスト + ページ番号（salescore-pptx-format 実測値）"""
    add_rect(slide, CONTENT_X, FOOTER_LINE_Y, CONTENT_RW, FOOTER_LINE_H, COLOR_PRIMARY)

    tb = slide.shapes.add_textbox(
        CONTENT_X, FOOTER_TEXT_Y, Inches(7.0), FOOTER_TEXT_H
    )
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = FOOTER_TEXT
    set_font(r, FONT_FOOTER, color=COLOR_MUTED)

    if page_num is not None:
        tb2 = slide.shapes.add_textbox(
            Inches(11.5), FOOTER_TEXT_Y, Inches(1.5), FOOTER_TEXT_H
        )
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        r2 = p2.add_run()
        r2.text = str(page_num)
        set_font(r2, FONT_PAGE_NUM, color=COLOR_MUTED)


# ─────────────────────────────────────────────
#  コンテンツレンダラー（共通ヘルパー）
# ─────────────────────────────────────────────

def render_text_items(slide, items, x, y, w, h):
    """
    bullet / sub_bullet / numbered / text / divider の混在リストを描画。
    divider で分割し、各テキストブロックに高さを等分配する。
    フォントサイズは salescore-pptx-format 準拠（16pt/14pt）。
    """
    blocks = []
    current_block = []
    for item in items:
        if item.get('type') == 'divider':
            if current_block:
                blocks.append(current_block)
                current_block = []
            blocks.append(None)
        else:
            current_block.append(item)
    if current_block:
        blocks.append(current_block)

    DIVIDER_H = Pt(12)
    n_dividers = blocks.count(None)
    n_text_blocks = len([b for b in blocks if b is not None])
    block_h = (h - n_dividers * DIVIDER_H) / max(n_text_blocks, 1) if n_text_blocks else h

    block_y = y
    for block in blocks:
        if block is None:
            add_rect(slide, x, block_y + Pt(4), w, Pt(1.0), COLOR_MID_GRAY)
            block_y += DIVIDER_H
        else:
            tb = slide.shapes.add_textbox(x, block_y, w, block_h)
            tf = tb.text_frame
            tf.word_wrap = True

            first = True
            for item in block:
                if first:
                    p = tf.paragraphs[0]
                    first = False
                else:
                    p = tf.add_paragraph()

                itype = item.get('type', 'bullet')
                text  = item.get('text', '')
                num   = item.get('num', '')

                if itype == 'bullet':
                    p.space_before = Pt(6)
                    p.space_after  = Pt(2)
                    r_sym = p.add_run()
                    r_sym.text = "● "
                    set_font(r_sym, FONT_BODY, color=COLOR_ACCENT)
                    add_inline_text(p, text, FONT_BODY, default_color=COLOR_PRIMARY)

                elif itype == 'sub_bullet':
                    p.space_before = Pt(3)
                    p.space_after  = Pt(1)
                    r_sym = p.add_run()
                    r_sym.text = "      – "
                    set_font(r_sym, FONT_BODY_SUB, color=COLOR_MUTED)
                    add_inline_text(p, text, FONT_BODY_SUB, default_color=COLOR_PRIMARY)

                elif itype == 'numbered':
                    p.space_before = Pt(6)
                    p.space_after  = Pt(2)
                    r_sym = p.add_run()
                    r_sym.text = f"{num}. "
                    set_font(r_sym, FONT_BODY, bold=True, color=COLOR_ACCENT)
                    add_inline_text(p, text, FONT_BODY, default_color=COLOR_PRIMARY)

                else:  # plain text
                    p.space_before = Pt(4)
                    p.space_after  = Pt(2)
                    add_inline_text(p, text, FONT_BODY, default_color=COLOR_PRIMARY)

            block_y += block_h


def render_table(slide, item, x, y, w, available_h=None):
    """
    テーブルアイテムをpptxテーブルとして描画。
    ヘッダー行: ダークネイビー背景・白太字・中央揃え
    データ行: ストライプ（薄グレー / 白）、word_wrap有効
    available_h指定時は行高さをスケールして空白を埋める
    戻り値: テーブルの高さ
    """
    headers = item.get('headers', [])
    rows    = item.get('rows', [])
    n_cols  = len(headers)
    if n_cols == 0:
        return Inches(0)

    MAX_ROW_H = int(Inches(1.2))
    n_data = len(rows)
    if available_h and n_data > 0:
        row_h = max(int(TABLE_DATA_ROW_H), min(MAX_ROW_H, (int(available_h) - int(TABLE_HEADER_ROW_H)) // n_data))
    else:
        row_h = int(TABLE_DATA_ROW_H)

    table_h = int(TABLE_HEADER_ROW_H) + row_h * n_data

    n_rows    = n_data + 1
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, table_h)
    tbl = tbl_shape.table

    # 列幅を均等に
    col_w = w // n_cols
    for j in range(n_cols):
        tbl.columns[j].width = col_w

    # ヘッダー行
    tbl.rows[0].height = TABLE_HEADER_ROW_H
    for j, hdr in enumerate(headers):
        cell = tbl.cell(0, j)
        set_cell_bg(cell, COLOR_DARK_NAVY)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_inline_text(p, hdr, FONT_BODY_SUB, default_bold=True, default_color=COLOR_WHITE)

    # データ行（ストライプ）
    for i, row_data in enumerate(rows):
        tbl.rows[i + 1].height = int(row_h)
        bg = COLOR_STRIPE if i % 2 == 0 else COLOR_WHITE
        for j in range(n_cols):
            cell = tbl.cell(i + 1, j)
            set_cell_bg(cell, bg)
            cell_text = row_data[j] if j < len(row_data) else ''
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            add_inline_text(p, cell_text, FONT_BODY_SUB, default_color=COLOR_PRIMARY)

    return table_h


def render_card_grid(slide, item, x, y, w, h):
    """
    カードグリッドを描画。
    N枚のカードを横に並べる。各カードにはカラー帯・アイコン・タイトル・本文。
    """
    cards = item.get('cards', [])
    if not cards:
        return
    N = len(cards)
    gap = int(Inches(0.25))
    card_w = (int(w) - gap * (N - 1)) // N
    card_h = int(h)

    for idx, card in enumerate(cards):
        cx = int(x) + idx * (card_w + gap)

        # 背景（薄グレー）
        bg = add_rect(slide, cx, int(y), card_w, card_h, COLOR_LIGHT_GRAY)
        remove_shadow(bg)

        # 上部アクセントストリップ
        strip_h = int(Pt(8))
        strip = add_rect(slide, cx, int(y), card_w, strip_h, card['color'])
        remove_shadow(strip)

        # アイコン（円）
        icon_d = int(Inches(0.5))
        icon_x = cx + (card_w - icon_d) // 2
        icon_y = int(y) + strip_h + int(Inches(0.15))
        oval = slide.shapes.add_shape(9, icon_x, icon_y, icon_d, icon_d)
        oval.fill.solid()
        oval.fill.fore_color.rgb = card['color']
        oval.line.fill.background()
        remove_shadow(oval)

        # アイコンラベル（白・太字・10pt）
        tf_oval = oval.text_frame
        tf_oval.word_wrap = False
        p_oval = tf_oval.paragraphs[0]
        p_oval.alignment = PP_ALIGN.CENTER
        r_oval = p_oval.add_run()
        r_oval.text = card.get('label', '')
        set_font(r_oval, 10, bold=True, color=COLOR_WHITE)

        # タイトルテキスト
        title_y = icon_y + icon_d + int(Inches(0.1))
        title_h = int(Inches(0.5))
        padding = int(Inches(0.12))
        tb_title = slide.shapes.add_textbox(cx + padding, title_y, card_w - padding * 2, title_h)
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.alignment = PP_ALIGN.CENTER
        r_title = p_title.add_run()
        r_title.text = card.get('title', '')
        set_font(r_title, 14, bold=True, color=COLOR_PRIMARY)
        remove_shadow(tb_title)

        # 本文テキスト
        body_y = title_y + title_h
        body_h = int(y) + card_h - body_y
        tb_body = slide.shapes.add_textbox(cx + padding, body_y, card_w - padding * 2, body_h)
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        p_body.alignment = PP_ALIGN.CENTER
        r_body = p_body.add_run()
        r_body.text = card.get('body', '')
        set_font(r_body, 12, color=COLOR_MUTED)
        remove_shadow(tb_body)


def render_banner(slide, item, x, y, w):
    """
    バナーを描画（ダークネイビー背景・白テキスト・14pt中央揃え）。
    戻り値: バナーの高さ（int EMU）
    """
    banner_h = int(Inches(0.55))
    bg = add_rect(slide, int(x), int(y), int(w), banner_h, COLOR_DARK_NAVY)
    remove_shadow(bg)

    tb = slide.shapes.add_textbox(int(x) + int(Inches(0.2)), int(y), int(w) - int(Inches(0.4)), banner_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_inline_text(p, item.get('text', ''), 14, default_color=COLOR_WHITE)
    remove_shadow(tb)

    return banner_h


def render_columns(slide, item, x, y, w, h):
    """
    2カラムレイアウトを描画。
    左右のカラムタイトル（任意）＋各カラムのコンテンツ＋中央縦区切り線。
    戻り値: 使用した高さ
    """
    GAP   = Inches(0.35)
    col_w = (w - GAP) / 2
    left_x  = x
    right_x = x + col_w + GAP

    title_h = Inches(0.38) if (item.get('left_title') or item.get('right_title')) else Inches(0)

    # カラムタイトル
    for col_x, title_text in ((left_x, item.get('left_title')), (right_x, item.get('right_title'))):
        if title_text:
            tb = slide.shapes.add_textbox(col_x, y, col_w, Inches(0.35))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            add_inline_text(p, title_text, FONT_BODY_SUB, default_bold=True, default_color=COLOR_DARK_NAVY)

    # 中央の縦区切り線
    div_x = x + col_w + GAP / 2 - Pt(0.5)
    add_rect(slide, div_x, y, Pt(1.0), h, COLOR_GRAY_LIGHT)

    content_y = y + title_h
    content_h = h - title_h

    if item.get('left'):
        render_text_items(slide, item['left'], left_x, content_y, col_w, content_h)
    if item.get('right'):
        render_text_items(slide, item['right'], right_x, content_y, col_w, content_h)

    return h


# ─────────────────────────────────────────────
#  スライド生成
# ─────────────────────────────────────────────

def make_title_slide(prs, title, subtitle, logo_path, page_num=1):
    """
    タイトルスライド
    ・上部 58%: ダークネイビー（ロゴ・タイトル）
    ・下部 42%: 白（サブタイトル・メタ情報）
    ・境界: ネイビーブルーのアクセントライン
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    split_y = SLIDE_HEIGHT * 0.58

    add_rect(slide, 0, 0, SLIDE_WIDTH, split_y, COLOR_DARK_NAVY)
    add_rect(slide, 0, split_y, SLIDE_WIDTH, SLIDE_HEIGHT - split_y, COLOR_WHITE)
    add_rect(slide, 0, split_y - Pt(3), SLIDE_WIDTH, Pt(6), COLOR_NAVY)

    title_y = split_y * 0.28
    tb = slide.shapes.add_textbox(
        Inches(1.0), title_y, SLIDE_WIDTH - Inches(3.8), split_y * 0.55
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = title
    set_font(r, 34, bold=True, color=COLOR_WHITE)

    if subtitle:
        sub_y = split_y * 0.28 + split_y * 0.55
        tb2 = slide.shapes.add_textbox(
            Inches(1.0), sub_y, SLIDE_WIDTH - Inches(3.8), Inches(0.9)
        )
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle.strip()
        set_font(r2, FONT_BODY_SUB, color=RGBColor(0xCC, 0xD8, 0xF0))

    add_logo(slide, logo_path)
    add_footer(slide, page_num)
    return slide


def make_section_slide(prs, title, section_num, logo_path, page_num=None):
    """
    セクション区切りスライド
    ・全面ダークネイビー
    ・セクション番号（小）＋ タイトル（大・白）
    ・ブランドブルーのアクセントライン
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, COLOR_DARK_NAVY)

    center_y = SLIDE_HEIGHT * 0.42

    if section_num:
        tb_num = slide.shapes.add_textbox(
            Inches(1.2), center_y - Inches(0.55),
            SLIDE_WIDTH - Inches(2.4), Inches(0.35)
        )
        tf_num = tb_num.text_frame
        p_num = tf_num.paragraphs[0]
        r_num = p_num.add_run()
        r_num.text = section_num
        set_font(r_num, 11, bold=True, color=COLOR_ACCENT)

    tb = slide.shapes.add_textbox(
        Inches(1.2), center_y - Inches(0.2) if section_num else center_y - Inches(0.4),
        SLIDE_WIDTH - Inches(2.4), Inches(1.2)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 36, bold=True, color=COLOR_WHITE)

    line_y = center_y + Inches(0.85) if not section_num else center_y + Inches(1.05)
    add_rect(slide, Inches(1.2), line_y, Inches(2.5), Pt(3), COLOR_BLUE)

    add_logo(slide, logo_path)
    add_footer(slide, page_num)
    return slide


def make_content_slide(prs, title, section_num, items, logo_path, page_num=None):
    """
    コンテンツスライド
    ・純白背景
    ・節番号（小グレー、任意）＋ メインタイトル（大・黒）
    ・ブラックヘッダーライン
    ・リード文 / 箇条書き / 番号付き / テーブル / 2カラム / 区切り線
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, COLOR_WHITE)

    # ── タイトルエリア（salescore-pptx-format: y=0.276, 21pt bold）──
    if section_num:
        tb_sec = slide.shapes.add_textbox(
            TITLE_AREA_X, Inches(0.07), TITLE_AREA_W, Inches(0.20)
        )
        tf_sec = tb_sec.text_frame
        p_sec = tf_sec.paragraphs[0]
        r_sec = p_sec.add_run()
        r_sec.text = section_num
        set_font(r_sec, 9, bold=True, color=COLOR_MUTED)

    tb_title = slide.shapes.add_textbox(
        TITLE_AREA_X, TITLE_AREA_Y, TITLE_AREA_W, TITLE_AREA_H
    )
    tf_title = tb_title.text_frame
    tf_title.word_wrap = False
    p_title = tf_title.paragraphs[0]
    r_title = p_title.add_run()
    r_title.text = title
    set_font(r_title, FONT_SLIDE_TITLE, bold=True, color=COLOR_PRIMARY)

    # ── ヘッダーライン（salescore-pptx-format: y=0.689）──
    add_header_line(slide)

    # ── コンテンツエリア ──
    lead_items    = [i for i in items if i.get('type') == 'lead']
    content_items = [i for i in items if i.get('type') != 'lead']

    current_y = CONTENT_AREA_Y

    # リード文（salescore-pptx-format: y=0.700, 18pt）
    if lead_items:
        lead_text = '　'.join(i['text'] for i in lead_items)
        tb_lead = slide.shapes.add_textbox(
            CONTENT_X, LEAD_Y, CONTENT_RW, LEAD_H
        )
        tf_lead = tb_lead.text_frame
        tf_lead.word_wrap = True
        p_lead = tf_lead.paragraphs[0]
        add_inline_text(p_lead, lead_text, FONT_LEAD, default_color=COLOR_PRIMARY)

    # ── バナーを先に抽出 ──
    banner_items  = [i for i in content_items if i.get('type') == 'banner']
    content_items = [i for i in content_items if i.get('type') != 'banner']

    BANNER_H   = int(Inches(0.55))
    BANNER_GAP = int(Inches(0.08))
    banner_reserved = len(banner_items) * (BANNER_H + BANNER_GAP)

    # ── セグメント分割 ──
    # table / columns / cards は独立セグメント、その他はテキストグループにまとめる
    segments = []
    text_buffer = []
    for item in content_items:
        if item.get('type') in ('table', 'columns', 'cards'):
            if text_buffer:
                segments.append(('text', list(text_buffer)))
                text_buffer = []
            segments.append((item['type'], item))
        else:
            text_buffer.append(item)
    if text_buffer:
        segments.append(('text', text_buffer))

    available_h = int(FOOTER_LINE_Y - current_y - Inches(0.1)) - banner_reserved

    # セグメント集計
    # dividerのみのtextセグメントはvar扱いしない（高さを消費しない）
    def is_divider_only(seg_data):
        return isinstance(seg_data, list) and all(
            item.get('type') == 'divider' for item in seg_data
        )

    n_tables = sum(1 for t, _ in segments if t == 'table')
    n_var    = sum(
        1 for t, d in segments
        if t in ('text', 'columns', 'cards') and not is_divider_only(d)
    )

    MIN_VAR_H = int(Inches(0.7))   # テキスト/カラムセグメント1つの最小高さ
    MAX_ROW_H = int(Inches(1.2))   # テーブル行高さの上限
    DIV_SEG_H = int(Inches(0.18))  # divider-onlyセグメントの高さ

    gap_total = int(Inches(0.12)) * max(n_tables - 1, 0)
    min_var_budget = MIN_VAR_H * n_var

    if n_tables > 0:
        table_budget = available_h - min_var_budget - gap_total
        per_table_h  = max(int(Inches(1.0)), table_budget // n_tables)
    else:
        per_table_h = None

    # テキスト/カラムに残りを等分（整数除算）
    TABLE_GAP = int(Inches(0.12))
    if n_var > 0:
        if n_tables > 0 and per_table_h:
            # テーブル高さ + 各テーブル後のギャップ（最後のテーブルの後も含む）
            actual_table_used = sum(
                min(MAX_ROW_H * len(data.get('rows', [])) + int(TABLE_HEADER_ROW_H), per_table_h)
                for seg_type, data in segments if seg_type == 'table'
            ) + TABLE_GAP * n_tables
        else:
            actual_table_used = 0
        n_div_segs = sum(1 for t, d in segments if t == 'text' and is_divider_only(d))
        var_h = (available_h - actual_table_used - DIV_SEG_H * n_div_segs) // n_var
        var_h = max(var_h, MIN_VAR_H)
    else:
        var_h = available_h

    # ── セグメント描画 ──
    y = current_y
    for seg_type, seg_data in segments:
        if seg_type == 'text':
            if is_divider_only(seg_data):
                # divider-onlyは小さな固定高さで処理
                render_text_items(slide, seg_data, CONTENT_X, y, CONTENT_RW, DIV_SEG_H)
                y += DIV_SEG_H
            else:
                render_text_items(slide, seg_data, CONTENT_X, y, CONTENT_RW, var_h)
                y += var_h
        elif seg_type == 'table':
            tbl_avail = per_table_h if per_table_h else None
            h = render_table(slide, seg_data, CONTENT_X, y, CONTENT_RW, available_h=tbl_avail)
            y += int(h) + int(Inches(0.12))
        elif seg_type == 'columns':
            render_columns(slide, seg_data, CONTENT_X, y, CONTENT_RW, var_h)
            y += var_h
        elif seg_type == 'cards':
            render_card_grid(slide, seg_data, int(CONTENT_X), int(y), int(CONTENT_RW), int(var_h))
            y += var_h

    # ── バナー描画（フッターラインの直上から上方向へ）──
    if banner_items:
        banner_y = int(FOOTER_LINE_Y) - BANNER_GAP - BANNER_H
        for b_item in reversed(banner_items):
            render_banner(slide, b_item, CONTENT_X, banner_y, CONTENT_RW)
            banner_y -= (BANNER_H + BANNER_GAP)

    add_logo(slide, logo_path)
    add_footer(slide, page_num)
    return slide


# ─────────────────────────────────────────────
#  Markdown パーサー
# ─────────────────────────────────────────────

def parse_markdown(md_text):
    """
    スライド構造のリストを返す。
    各要素:
      title   : {type, title, subtitle}
      section : {type, title, section_num}
      content : {type, title, section_num, items}
    """
    slides        = []
    current       = None
    col_side      = None   # None | 'left' | 'right'
    table_header  = None   # 保留中のテーブルヘッダー行
    table_active  = None   # 処理中のテーブルアイテム
    card_mode     = False  # ::cards ブロック内かどうか
    card_cur      = None   # 現在処理中のカード辞書
    cards_list    = None   # 現在の::cardsブロックのカードリスト

    def flush_table():
        nonlocal table_header, table_active
        if table_active and current and current.get('type') == 'content':
            _append_item(table_active)
        table_header = None
        table_active = None

    def _append_item(item):
        """col_side に応じてアイテムを適切な場所に追加"""
        if (col_side
                and current.get('items')
                and current['items'][-1].get('type') == 'columns'):
            current['items'][-1][col_side].append(item)
        else:
            current['items'].append(item)

    for raw_line in md_text.splitlines():
        line = raw_line.strip()

        # ── テーブル行の検出（| で始まる行）──
        if line.startswith('|'):
            if re.match(r'^[\|\s\-:]+$', line):
                # 区切り行（|---|---|）→ ヘッダー確定
                if table_header is not None and current and current.get('type') == 'content':
                    table_active = {
                        'type': 'table',
                        'headers': table_header,
                        'rows': []
                    }
                    table_header = None
            else:
                # データ行
                cells = [c.strip() for c in line.strip('|').split('|')]
                if table_active is not None:
                    table_active['rows'].append(cells)
                else:
                    flush_table()
                    table_header = cells
            continue

        # テーブル以外の行が来たらテーブルを確定
        flush_table()

        # ── ::cards ブロック内の処理 ──
        if card_mode:
            if line == ':::':
                # カードブロック終了
                card_mode_local = True  # will reset below
                nonlocal_card_mode = False
                if card_cur is not None:
                    cards_list.append(card_cur)
                if current and current.get('type') == 'content':
                    current['items'].append({'type': 'cards', 'cards': cards_list})
                # reset card state in outer scope via reassignment trick
                # We use a mutable approach: set flags via the variables directly
                # (done after continue)
            elif re.match(r'^\[(.+?)\]\s*#([0-9A-Fa-f]{6})$', line):
                # 新しいカード開始
                if card_cur is not None:
                    cards_list.append(card_cur)
                m = re.match(r'^\[(.+?)\]\s*#([0-9A-Fa-f]{6})$', line)
                label = m.group(1)
                hex_color = m.group(2)
                r_val = int(hex_color[0:2], 16)
                g_val = int(hex_color[2:4], 16)
                b_val = int(hex_color[4:6], 16)
                card_cur = {
                    'label': label,
                    'color': RGBColor(r_val, g_val, b_val),
                    'title': None,
                    'body': '',
                }
                continue
            else:
                # タイトルまたは本文
                if card_cur is not None:
                    if card_cur['title'] is None:
                        card_cur['title'] = line
                    else:
                        if card_cur['body']:
                            card_cur['body'] += ' ' + line
                        else:
                            card_cur['body'] = line
                continue
            # Handle ':::' close for card_mode
            if line == ':::':
                card_mode = False
                card_cur = None
                cards_list = None
                continue

        # ── H3 → セクション区切りスライド ──
        if line.startswith('### '):
            if current:
                slides.append(current)
            raw_title = line[4:].strip()
            if ' | ' in raw_title:
                sec, ttl = raw_title.split(' | ', 1)
                current = {'type': 'section', 'title': ttl.strip(),
                           'section_num': sec.strip(), 'items': []}
            else:
                current = {'type': 'section', 'title': raw_title,
                           'section_num': None, 'items': []}
            col_side = None

        # ── H2 → コンテンツスライド ──
        elif line.startswith('## '):
            if current:
                slides.append(current)
            raw_title = line[3:].strip()
            if ' | ' in raw_title:
                sec, ttl = raw_title.split(' | ', 1)
                current = {'type': 'content', 'title': ttl.strip(),
                           'section_num': sec.strip(), 'items': []}
            else:
                current = {'type': 'content', 'title': raw_title,
                           'section_num': None, 'items': []}
            col_side = None

        # ── H1 → タイトルスライド ──
        elif line.startswith('# '):
            if current:
                slides.append(current)
            current = {
                'type': 'title',
                'title': line[2:].strip(),
                'subtitle': '',
                'items': [],
            }
            col_side = None

        # ── ::cards ブロック開始 ──
        elif line.startswith('::cards'):
            flush_table()
            card_mode  = True
            card_cur   = None
            cards_list = []

        # ── 2カラムレイアウト ──
        elif line.startswith('::left'):
            col_side  = 'left'
            col_title = line[6:].strip() or None
            if current and current.get('type') == 'content':
                if not current['items'] or current['items'][-1].get('type') != 'columns':
                    current['items'].append({
                        'type': 'columns',
                        'left': [], 'right': [],
                        'left_title': col_title, 'right_title': None
                    })
                elif col_title:
                    current['items'][-1]['left_title'] = col_title

        elif line.startswith('::right'):
            col_side  = 'right'
            col_title = line[7:].strip() or None
            if current and current.get('type') == 'content':
                if not current['items'] or current['items'][-1].get('type') != 'columns':
                    current['items'].append({
                        'type': 'columns',
                        'left': [], 'right': [],
                        'left_title': None, 'right_title': col_title
                    })
                elif col_title:
                    current['items'][-1]['right_title'] = col_title

        elif line == ':::':
            col_side = None

        # ── バナー（>> text）──
        elif line.startswith('>> '):
            flush_table()
            text = line[3:].strip()
            if current and current.get('type') == 'content':
                _append_item({'type': 'banner', 'text': text})

        # ── リード文（> blockquote）──
        elif line.startswith('> '):
            text = line[2:].strip()
            if current and current.get('type') == 'content':
                current['items'].append({'type': 'lead', 'text': text})

        # ── 区切り線 --- ──
        elif line == '---':
            if current and current.get('type') == 'content':
                _append_item({'type': 'divider'})

        # ── 字下げ箇条書き（インデント付き - or *）──
        elif re.match(r'^\s{2,}[\*\-] ', raw_line):
            text = re.sub(r'^\s+[\*\-] ', '', raw_line).strip()
            if current:
                _append_item({'type': 'sub_bullet', 'text': text})

        # ── 標準箇条書き（* または -）──
        elif re.match(r'^[\*\-] ', line):
            text = line[2:].strip()
            if current:
                _append_item({'type': 'bullet', 'text': text})

        # ── 番号付きリスト ──
        elif re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line).strip()
            # 同一コンテキスト内での連番
            if (col_side
                    and current.get('items')
                    and current['items'][-1].get('type') == 'columns'):
                col_items = current['items'][-1][col_side]
                num = len([i for i in col_items if i.get('type') == 'numbered']) + 1
            else:
                num = len([i for i in current.get('items', [])
                           if i.get('type') == 'numbered']) + 1
            if current:
                _append_item({'type': 'numbered', 'text': text, 'num': num})

        # ── タイトルスライドのサブタイトル行 ──
        elif line and current and current['type'] == 'title':
            current['subtitle'] += line + ' '

        # ── その他テキスト（コンテンツスライドの平文）──
        elif line and current and current.get('type') == 'content':
            _append_item({'type': 'text', 'text': line})

    flush_table()
    if current:
        slides.append(current)

    return slides


# ─────────────────────────────────────────────
#  メイン
# ─────────────────────────────────────────────

def generate_ppt(md_path, out_path, logo_path=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    slides_data = parse_markdown(md_text)

    prs = Presentation()
    prs.slide_width  = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    for i, s in enumerate(slides_data, start=1):
        stype = s['type']

        if stype == 'title':
            make_title_slide(
                prs, s['title'], s.get('subtitle', '').strip(),
                logo_path, page_num=i
            )
        elif stype == 'section':
            make_section_slide(
                prs, s['title'], s.get('section_num'),
                logo_path, page_num=i
            )
        else:  # content
            items = s.get('items', [])
            if items or s.get('title'):
                make_content_slide(
                    prs, s['title'], s.get('section_num'),
                    items, logo_path, page_num=i
                )
            else:
                make_section_slide(
                    prs, s['title'], s.get('section_num'),
                    logo_path, page_num=i
                )

    prs.save(out_path)
    print(f"✅ 保存完了: {out_path}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("使い方: python generate_ppt.py input.md output.pptx [logo.png]")
        sys.exit(1)

    md_file   = sys.argv[1]
    out_file  = sys.argv[2]
    logo_file = sys.argv[3] if len(sys.argv) >= 4 else None

    generate_ppt(md_file, out_file, logo_file)
